# 다수 기업 검색 API 추가 코드

@main.route('/api/multi-company-search', methods=['POST'])
@login_required
def api_multi_company_search():
    """다수 기업 일괄 검색 API (스트리밍)"""
    import zipfile
    from io import BytesIO
    
    try:
        data = request.get_json()
        companies = data.get('companies', [])
        start_date = data.get('start_date', '')
        end_date = data.get('end_date', '')
        
        if not companies:
            return jsonify({"error": "검색할 기업을 선택해주세요"}), 400
        
        if not start_date or not end_date:
            return jsonify({"error": "검색 기간을 설정해주세요"}), 400
        
        print(f"[DEBUG] 다수 기업 검색 시작: {len(companies)}개 기업, 기간: {start_date} ~ {end_date}")
        
        def generate():
            success_count = 0
            error_count = 0
            excel_files = []  # (filename, buffer) 튜플 리스트
            
            try:
                for idx, company in enumerate(companies, 1):
                    try:
                        print(f"[DEBUG] [{idx}/{len(companies)}] {company} 검색 시작")
                        
                        # 진행 상황 전송
                        progress_data = {
                            'type': 'progress',
                            'company': company,
                            'completed': idx - 1,
                            'total': len(companies),
                            'status': 'processing',
                            'message': '검색 중...'
                        }
                        yield f"data: {json.dumps(progress_data, ensure_ascii=False)}\n\n"
                        
                        # 뉴스 검색
                        news_items = news_service.search_news_by_period(
                            query=company,
                            start_date=start_date,
                            end_date=end_date
                        )
                        
                        print(f"[DEBUG] {company}: {len(news_items)}개 뉴스 검색됨")
                        
                        if len(news_items) == 0:
                            # 검색 결과 없음
                            progress_data = {
                                'type': 'progress',
                                'company': company,
                                'completed': idx,
                                'total': len(companies),
                                'status': 'success',
                                'message': '검색 결과 없음 (0개)'
                            }
                            yield f"data: {json.dumps(progress_data, ensure_ascii=False)}\n\n"
                            error_count += 1
                            continue
                        
                        # 엑셀 파일 생성
                        excel_buffer = create_simple_excel(company, news_items, start_date, end_date)
                        
                        # 파일명 생성
                        filename = f"{company}_뉴스검색결과_{start_date.replace('-', '')}_{end_date.replace('-', '')}.xlsx"
                        excel_files.append((filename, excel_buffer))
                        
                        success_count += 1
                        
                        # 성공 진행 상황 전송
                        progress_data = {
                            'type': 'progress',
                            'company': company,
                            'completed': idx,
                            'total': len(companies),
                            'status': 'success',
                            'message': f'완료 ({len(news_items)}개 뉴스)'
                        }
                        yield f"data: {json.dumps(progress_data, ensure_ascii=False)}\n\n"
                        
                    except Exception as e:
                        print(f"[ERROR] {company} 검색 실패: {str(e)}")
                        error_count += 1
                        
                        # 오류 진행 상황 전송
                        progress_data = {
                            'type': 'progress',
                            'company': company,
                            'completed': idx,
                            'total': len(companies),
                            'status': 'error',
                            'message': f'오류: {str(e)}'
                        }
                        yield f"data: {json.dumps(progress_data, ensure_ascii=False)}\n\n"
                
                # 모든 검색 완료 - ZIP 파일 생성
                if excel_files:
                    print(f"[DEBUG] ZIP 파일 생성 시작: {len(excel_files)}개 파일")
                    
                    # ZIP 파일 생성
                    zip_buffer = BytesIO()
                    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
                        for filename, excel_buffer in excel_files:
                            zip_file.writestr(filename, excel_buffer.getvalue())
                    
                    # ZIP 파일 저장
                    zip_filename = f"다수기업검색결과_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
                    zip_path = os.path.join(tempfile.gettempdir(), zip_filename)
                    
                    with open(zip_path, 'wb') as f:
                        f.write(zip_buffer.getvalue())
                    
                    print(f"[DEBUG] ZIP 파일 저장 완료: {zip_path}")
                    
                    # 완료 메시지 전송
                    complete_data = {
                        'type': 'complete',
                        'success_count': success_count,
                        'error_count': error_count,
                        'total': len(companies),
                        'zip_filename': zip_filename
                    }
                    yield f"data: {json.dumps(complete_data, ensure_ascii=False)}\n\n"
                else:
                    # 모든 검색 실패
                    error_data = {
                        'type': 'error',
                        'message': '모든 기업 검색에 실패했습니다.'
                    }
                    yield f"data: {json.dumps(error_data, ensure_ascii=False)}\n\n"
                    
            except Exception as e:
                print(f"[ERROR] 다수 기업 검색 오류: {str(e)}")
                import traceback
                print(f"[ERROR] 오류 상세: {traceback.format_exc()}")
                
                error_data = {
                    'type': 'error',
                    'message': f'검색 중 오류가 발생했습니다: {str(e)}'
                }
                yield f"data: {json.dumps(error_data, ensure_ascii=False)}\n\n"
        
        response = Response(generate(), mimetype='text/event-stream')
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
        response.headers['X-Accel-Buffering'] = 'no'
        response.headers['Connection'] = 'keep-alive'
        return response
        
    except Exception as e:
        print(f"[ERROR] 다수 기업 검색 API 오류: {str(e)}")
        return jsonify({"error": f"검색 요청 처리 중 오류가 발생했습니다: {str(e)}"}), 500


@main.route('/api/download-multi-result/<filename>')
@login_required
def api_download_multi_result(filename):
    """다수 기업 검색 결과 ZIP 파일 다운로드"""
    try:
        zip_path = os.path.join(tempfile.gettempdir(), filename)
        
        if not os.path.exists(zip_path):
            return jsonify({"error": "파일을 찾을 수 없습니다."}), 404
        
        # 파일 전송 후 삭제
        response = send_file(
            zip_path,
            as_attachment=True,
            download_name=filename,
            mimetype='application/zip'
        )
        
        # 파일 전송 후 삭제 (백그라운드)
        def delete_file():
            import time
            time.sleep(5)  # 5초 대기 후 삭제
            try:
                if os.path.exists(zip_path):
                    os.remove(zip_path)
                    print(f"[DEBUG] 임시 ZIP 파일 삭제: {zip_path}")
            except Exception as e:
                print(f"[ERROR] 임시 ZIP 파일 삭제 실패: {str(e)}")
        
        thread = threading.Thread(target=delete_file)
        thread.daemon = True
        thread.start()
        
        return response
        
    except Exception as e:
        print(f"[ERROR] ZIP 파일 다운로드 오류: {str(e)}")
        return jsonify({"error": str(e)}), 500


def create_simple_excel(company_name, news_items, start_date, end_date):
    """간단한 뉴스 검색 결과 엑셀 파일 생성"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from io import BytesIO
    
    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스 검색 결과"
    
    # 헤더 스타일
    header_fill = PatternFill(start_color="4285f4", end_color="4285f4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # 테두리 스타일
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 제목 행
    ws.merge_cells('A1:E1')
    title_cell = ws['A1']
    title_cell.value = f"{company_name} 뉴스 검색 결과 ({start_date} ~ {end_date})"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
    ws.row_dimensions[1].height = 30
    
    # 헤더
    headers = ["번호", "제목", "내용", "출처", "날짜"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    ws.row_dimensions[2].height = 25
    
    # 데이터 행
    for idx, news in enumerate(news_items, 1):
        row = idx + 2
        
        ws.cell(row=row, column=1).value = idx
        ws.cell(row=row, column=2).value = news.get('title', '')
        ws.cell(row=row, column=3).value = news.get('description', '')
        ws.cell(row=row, column=4).value = news.get('source', '')
        ws.cell(row=row, column=5).value = news.get('published', '')
        
        # 스타일 적용
        for col in range(1, 6):
            cell = ws.cell(row=row, column=col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border
    
    # 열 너비 조정
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 60
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 20
    
    # BytesIO 버퍼에 저장
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    return excel_buffer
